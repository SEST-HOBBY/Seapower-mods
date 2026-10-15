#!/usr/bin/env python3
import csv
from pathlib import Path
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.formatting.rule import CellIsRule

DIR=Path(__file__).resolve().parent

def data(name):
    with (DIR/name).open(encoding='utf-8-sig',newline='') as f:
        return list(csv.DictReader(f))

master=data('SEST_IMAGE_MASTER_20261014.csv')
pins=data('PINNED_VERIFIED_FILE_PAGES.csv')
anduril=data('ANDURIL_VARIANT_RECONCILIATION.csv')
wb=Workbook();summary=wb.active;summary.title='START HERE'
navy='112334';green='178F74';white='FFFFFF';cyan='76D3D5';fog='EEF3F7';review='FFF2D6';hold='FFE8E8';grey='667789';black='18222C';blue='206AA5'

summary.sheet_view.showGridLines=False
summary.column_dimensions['A'].width=5
summary.column_dimensions['B'].width=39
summary.column_dimensions['C'].width=34
summary.column_dimensions['D'].width=31
summary.merge_cells('B2:D3')
cell=summary['B2'];cell.value='SEST / SEA POWER\nIMAGE ACQUISITION MASTER';cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(name='Aptos Display',size=20,bold=True,color=white);cell.alignment=Alignment(wrap_text=True,vertical='center')
summary.row_dimensions[2].height=28;summary.row_dimensions[3].height=30
rows=[
 ('Audit snapshot','14 October 2026','Not yet merged into Sea Power'),
 ('Legacy actionable gaps',89,'Gallery prioritised missing/weak images'),
 ('October addition/family records',27,'Includes pending Anduril source family IDs'),
 ('Total mapping actions',len(master),'Not equal to number of downloaded photos'),
 ('Pinned Commons File pages',len(pins),'12 licences checked online; all rechecked at download'),
 ('Commons category source queues',sum(x['acquisition_mode']=='category' for x in master),'Human visual approval required'),
 ('Commons search source queues',sum(x['acquisition_mode']=='search' for x in master),'Human visual approval required'),
 ('Anduril permission holds',sum(x['acquisition_mode']=='permission' for x in master),'Official/author images require rights clearance'),
 ('Actual photo downloads in this ZIP',0,'Run fetch_images.py on a networked Windows PC'),
 ('','', ''),
 ('Next step','Run RUN_WINDOWS.cmd','Then inspect CONTACT_SHEET.html'),
 ('Release rule','No AI images','Keep original creator and licence credits'),
 ('Repo guard','Read-only first','No direct modification or pushes to main'),
]
for j,row in enumerate(rows,start=6):
    for k,val in enumerate(row,start=2):
        c=summary.cell(j,k,val);c.alignment=Alignment(vertical='center',wrap_text=True)
        c.font=Font(name='Aptos',size=11,color=black,bold=(k==2))
        if j%2==0:c.fill=PatternFill('solid',fgColor=fog)
    summary.row_dimensions[j].height=28
summary.freeze_panes='B6'


def add_sheet(name, rows, extra=False):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=80
    keys=list(rows[0].keys()) if rows else []
    if extra: keys+=['reviewer_decision','reviewer_note']
    for i,k in enumerate(keys,1):
        c=ws.cell(1,i,k.replace('_',' ').upper());c.font=Font(name='Aptos',bold=True,color=white,size=10)
        c.fill=PatternFill('solid',fgColor=navy);c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=40
    for rno,row in enumerate(rows,2):
        mode=row.get('acquisition_mode','')
        for cno,k in enumerate(keys,1):
            value=row.get(k,'')
            c=ws.cell(rno,cno,value);c.font=Font(name='Aptos',size=10,color=green if k in ('source_page','source_ref','licence_checked_from_web') else black)
            c.alignment=Alignment(vertical='center',wrap_text=k in ('unit_label','notes','import_status'))
            if rno%2==0:c.fill=PatternFill('solid',fgColor='F6FAFB')
            if k in ('source_page','source_ref') and isinstance(value,str) and value.startswith('https://'):
                c.hyperlink=value;c.font=Font(name='Aptos',size=10,color=blue,underline='single')
            if k=='acquisition_mode':
                if mode=='file':c.fill=PatternFill('solid',fgColor='DEEEE6')
                elif mode in ('category','search'):c.fill=PatternFill('solid',fgColor=review)
                else:c.fill=PatternFill('solid',fgColor=hold)
            if k in ('reviewer_decision','reviewer_note'):c.font=Font(name='Aptos',size=10,color=blue)
        ws.row_dimensions[rno].height=35 if name=='MASTER' else 26
    widths={'unit_id':29,'unit_label':38,'origin':34,'provider_id':22,'was_missing':18,'legacy_missions':16,
     'priority':34,'asset_group':30,'acquisition_mode':19,'source_ref':78,'source_page':73,
     'image_exactness':28,'licence_checked_from_web':23,'import_status':65,'notes':100,
     'filename':78,'web_verified_licence':23,'reviewer_decision':22,'reviewer_note':42,
     'workshop_mod':18,'family':27,'unit_id_from_export':30,'variant_name_from_export':35,'shared_asset':30,'status':60}
    for i,k in enumerate(keys,1):ws.column_dimensions[get_column_letter(i)].width=widths.get(k,27)
    ws.freeze_panes='D2' if name=='MASTER' else 'B2'
    ws.auto_filter.ref=f'A1:{get_column_letter(len(keys))}{len(rows)+1}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws

add_sheet('MASTER',master,extra=True)
add_sheet('PINNED FILE PAGES',pins)
add_sheet('ANDURIL IDs TO FILL',anduril,extra=True)
wb.active=0
wb.save(DIR/'SEST_Image_Master_20261014.xlsx')
print('Wrote Excel workbook',len(master),'rows')
